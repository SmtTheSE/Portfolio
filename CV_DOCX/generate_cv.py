#!/usr/bin/env python3
"""Generate Sitt Min Thar CV (DOCX + PDF) in the classic one-column
Experience/Education/Projects/Skills academic-CV template style."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fpdf import FPDF

OUT_DIR = Path(__file__).resolve().parent
DOCX_PATH = OUT_DIR / "Sitt_Min_Thar_Resume.docx"
PDF_PATH = OUT_DIR / "Sitt_Min_Thar_Resume.pdf"

FONT = "Times New Roman"
LINK_BLUE = RGBColor(0x05, 0x63, 0xC1)
BLACK = RGBColor(0x00, 0x00, 0x00)

NAME = "SITT MIN THAR"
EMAIL = "sittminthar005@gmail.com"
GITHUB = "github.com/SmtTheSE"
WEBSITE = "sittminthar.from-mm.dev"

EXPERIENCE = [
    {
        "title": "Software Engineer",
        "org": "Ad Venture Studio",
        "dates": "2026 -- Present",
        "bullets": [
            "Building and shipping mobile apps for App Store and Google Play, owning product engineering, store releases, and app infrastructure",
            "Partnering with design and monetization leads on UX, performance, and retention-oriented product engineering",
        ],
    },
    {
        "title": "Founding Engineer",
        "org": "Bravestep",
        "org_url": "https://bravestep.ai/",
        "dates": "Present",
        "bullets": [
            "Founding engineering role focused on systems direction and product infrastructure",
        ],
    },
    {
        "title": "Release-Acceptance Engineer",
        "org": "Malibu",
        "org_url": "https://malibu.tech/",
        "dates": "Present",
        "bullets": [
            "Independent release-acceptance engineer for a macOS distributed-compute marketplace, serving as the physical-hardware validation gate between CI-green and fleet rollout",
        ],
    },
    {
        "title": "Software Development Intern",
        "org": "AIOT Inc",
        "org_url": "https://www.facebook.com/aiot.global.inc",
        "dates": "Internship",
        "bullets": [
            "Internship focused on backend integration, application structure, and production-minded software practices",
        ],
    },
]

EDUCATION = [
    {
        "degree": "B.S. Data Science",
        "school": "Saigon Business School, Vietnam",
        "dates": "Expected 2027",
    },
    {
        "degree": "B.S. Software Engineering",
        "school": "University of Computer Studies, Yangon",
        "dates": "2022 -- 2024",
    },
]

PROJECTS = [
    {
        "name": "Rental Property AI Platform",
        "url": "https://github.com/SmtTheSE/Rental-Property-AI-Platform",
        "bullets": [
            "10M-dataset-powered demand forecasting and gap analysis engine processing rental listings across 40 Indian cities (Python, Flask, Scikit-learn, Next.js)",
        ],
    },
    {
        "name": "Ultimate ByteMe: Pan-Asian Real Estate Intelligence",
        "url": "https://github.com/SmtTheSE/Conference",
        "bullets": [
            "Computational intelligence platform for the 1st Synergia International Conference 2026 with a LightGBM valuation engine, an investment opportunity scanner, and a RAG-grounded AI assistant on Qwen2.5/Gemini (Next.js, Python, Flask, LightGBM, Ollama)",
        ],
    },
    {
        "name": "SBS Student Serving System",
        "url": "https://github.com/SmtTheSE/docker_SBS",
        "bullets": [
            "Full-stack academic platform for Saigon Business School with a Spring Boot 3 REST API and React 18 frontend, JWT authentication, and role-based access control (Java 21, Spring Boot 3, React 18, Docker, AWS S3)",
        ],
    },
    {
        "name": "DataCleanr",
        "url": "https://github.com/SmtTheSE/Data-Cleanr",
        "bullets": [
            "AI-powered dataset cleaning app with industry detection across 12 sectors and automated CSV/Excel harmonization (Python, Cython, JavaScript)",
        ],
    },
    {
        "name": "DentalBridge",
        "url": "https://github.com/SmtTheSE/DentalBridge",
        "bullets": [
            "AI treatment coordinator using Google Gemini 1.5 Pro to translate clinical terminology into patient-friendly guidance (FastAPI, Next.js)",
        ],
    },
    {
        "name": "Kamisori",
        "url": "https://github.com/SmtTheSE/Kamisori",
        "bullets": [
            "“Embrace the Future” -- e-commerce platform for a local clothing brand in Myanmar; architected a secure, scalable backend ensuring data integrity (PostgreSQL, Supabase, Edge Functions)",
        ],
    },
    {
        "name": "Dental Blinding & Age Estimation",
        "url": "https://github.com/SmtTheSE/Dental-Blinding-Process-Project",
        "bullets": [
            "Internal pediatric dental age estimation system for the University of Dental Medicine, Mandalay, implementing the AlQahtani and Demirjian methods with a role-based Supervisor/PI blinding workflow to eliminate analysis bias (Python, Flask, PostgreSQL)",
        ],
    },
]

OPEN_SOURCE = [
    {
        "name": "BuddyUsage",
        "url": "https://github.com/SmtTheSE/BuddyUsage",
        "bullets": [
            "Menu-bar usage tracker for Claude, Codex, Gemini, Cursor and Copilot with pace prediction and local agent control; exposes usage as a local MCP server. 25 releases, macOS/Windows/Linux (Electron, TypeScript)",
        ],
    },
]

SKILLS = [
    ("Programming", "Python (Flask, FastAPI, Pandas, NumPy, Scikit-learn, LightGBM), Java (Spring Boot), Go, PHP, JavaScript/TypeScript, SQL"),
    ("Software", "Docker, Nginx, AWS (S3, EC2), PostgreSQL, MySQL, Git, GitHub, VPS/SSL, CI/CD"),
]


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

def add_bottom_rule(paragraph, size: int = 6, color: str = "000000") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def section_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.name = FONT
    run.font.size = Pt(11.5)
    run.font.color.rgb = BLACK
    add_bottom_rule(p)


def add_hyperlink(paragraph, text: str, url: str, bold: bool = True, size: float = 11) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run_el = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")

    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), FONT)
    rfonts.set(qn("w:hAnsi"), FONT)
    rpr.append(rfonts)

    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)

    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)

    if bold:
        rpr.append(OxmlElement("w:b"))

    run_el.append(rpr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run_el.append(text_el)
    hyperlink.append(run_el)
    paragraph._p.append(hyperlink)


def two_col_line(doc: Document, left: str, right: str, bold: bool = True, size: float = 10.5, left_url: str | None = None, left_link_text: str | None = None) -> None:
    left_w, right_w = Inches(6.2), Inches(0.8)
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.allow_autofit = False

    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    tbl_w = OxmlElement("w:tblW")
    tbl_w.set(qn("w:w"), str(int(left_w.twips + right_w.twips)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_w)

    left_cell, right_cell = table.rows[0].cells
    left_cell.width = left_w
    right_cell.width = right_w
    for cell, width in ((left_cell, left_w), (right_cell, right_w)):
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_w = OxmlElement("w:tcW")
        tc_w.set(qn("w:w"), str(int(width.twips)))
        tc_w.set(qn("w:type"), "dxa")
        tc_pr.append(tc_w)
        tc_mar = OxmlElement("w:tcMar")
        for side in ("left", "right"):
            m = OxmlElement(f"w:{side}")
            m.set(qn("w:w"), "0")
            m.set(qn("w:type"), "dxa")
            tc_mar.append(m)
        tc_pr.append(tc_mar)

    lp = left_cell.paragraphs[0]
    lp.paragraph_format.space_after = Pt(0)
    if left_url and left_link_text:
        prefix, _, suffix = left.partition(left_link_text)
        if prefix:
            pr = lp.add_run(prefix)
            pr.bold = bold
            pr.font.name = FONT
            pr.font.size = Pt(size)
        add_hyperlink(lp, left_link_text, left_url, bold=bold, size=size)
        if suffix:
            sr = lp.add_run(suffix)
            sr.bold = bold
            sr.font.name = FONT
            sr.font.size = Pt(size)
    else:
        lr = lp.add_run(left)
        lr.bold = bold
        lr.font.name = FONT
        lr.font.size = Pt(size)

    rp = right_cell.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rp.paragraph_format.space_after = Pt(0)
    rr = rp.add_run(right)
    rr.bold = bold
    rr.font.name = FONT
    rr.font.size = Pt(size)
    return table


def add_link_run(paragraph, text: str, bold: bool = True, size: float = 11) -> None:
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = LINK_BLUE
    run.font.underline = True


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.25)
        run = p.add_run(item)
        run.font.name = FONT
        run.font.size = Pt(10.5)


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    list_bullet = doc.styles["List Bullet"]
    list_bullet.font.name = FONT
    list_bullet.paragraph_format.space_before = Pt(0)
    list_bullet.paragraph_format.space_after = Pt(2)
    list_bullet.paragraph_format.line_spacing = 1.0

    # --- Name ---
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_p.paragraph_format.space_after = Pt(4)
    nr = name_p.add_run(NAME)
    nr.bold = False
    nr.font.name = FONT
    nr.font.size = Pt(22)
    nr.font.color.rgb = BLACK

    # --- Contact line ---
    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_p.paragraph_format.space_after = Pt(6)
    add_link_run(contact_p, GITHUB, bold=False, size=10.5)
    sep = contact_p.add_run("  |  ")
    sep.font.name = FONT
    sep.font.size = Pt(10.5)
    add_link_run(contact_p, EMAIL, bold=False, size=10.5)
    sep2 = contact_p.add_run("  |  ")
    sep2.font.name = FONT
    sep2.font.size = Pt(10.5)
    add_link_run(contact_p, WEBSITE, bold=False, size=10.5)

    # --- Experience ---
    section_heading(doc, "Experience")
    for exp in EXPERIENCE:
        left = f"{exp['title']} | {exp['org']}"
        org_url = exp.get("org_url")
        two_col_line(doc, left, exp["dates"], left_url=org_url, left_link_text=exp["org"] if org_url else None)
        add_bullets(doc, exp["bullets"])

    # --- Education ---
    section_heading(doc, "Education")
    for edu in EDUCATION:
        two_col_line(doc, f"{edu['degree']} -- {edu['school']}", edu["dates"])

    # --- Projects ---
    section_heading(doc, "Projects")
    for proj in PROJECTS:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        add_hyperlink(p, proj["name"], proj["url"])
        add_bullets(doc, proj["bullets"])

    # --- Open Source ---
    section_heading(doc, "Open Source")
    for proj in OPEN_SOURCE:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        add_hyperlink(p, proj["name"], proj["url"])
        add_bullets(doc, proj["bullets"])

    # --- Skills ---
    section_heading(doc, "Skills")
    for label, value in SKILLS:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        lr = p.add_run(f"{label}: ")
        lr.bold = True
        lr.font.name = FONT
        lr.font.size = Pt(10.5)
        vr = p.add_run(value)
        vr.font.name = FONT
        vr.font.size = Pt(10.5)

    doc.save(DOCX_PATH)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def pdf_text(text: str) -> str:
    """fpdf core fonts are latin-1 only."""
    return (
        text.replace("–", "-")
        .replace("—", "--")
        .replace("’", "'")
        .replace("“", '"')
        .replace("”", '"')
    )


class CVPdf(FPDF):
    def __init__(self):
        super().__init__()
        self.set_auto_page_break(auto=True, margin=5)
        self.set_margins(15, 7, 15)
        self.link_blue = (5, 99, 193)

    def section_heading(self, title: str) -> None:
        self.ln(1.5)
        self.set_font("Times", "B", 12.5)
        self.set_text_color(0, 0, 0)
        self.cell(0, 6, pdf_text(title.upper()), new_x="LMARGIN", new_y="NEXT")
        y = self.get_y()
        self.set_draw_color(0, 0, 0)
        self.line(self.l_margin, y, self.w - self.r_margin, y)
        self.ln(1.5)

    def two_col_line(self, left: str, right: str, size: float = 10.5) -> None:
        self.set_font("Times", "B", size)
        self.set_text_color(0, 0, 0)
        usable = self.w - self.l_margin - self.r_margin
        right_w = 40.0
        left_w = usable - right_w
        self.cell(left_w, 5.6, pdf_text(left))
        self.cell(right_w, 5.6, pdf_text(right), align="R", new_x="LMARGIN", new_y="NEXT")

    def exp_line(self, title: str, org: str, right: str, size: float = 10.5, org_url: str | None = None) -> None:
        self.set_font("Times", "B", size)
        self.set_text_color(0, 0, 0)
        usable = self.w - self.l_margin - self.r_margin
        right_w = 40.0
        left_w = usable - right_w

        x0, y0 = self.get_x(), self.get_y()
        self.write(5.6, pdf_text(f"{title} | "))
        if org_url:
            self.set_text_color(*self.link_blue)
            self.write(5.6, pdf_text(org), link=org_url)
            self.set_text_color(0, 0, 0)
        else:
            self.write(5.6, pdf_text(org))

        self.set_xy(x0 + left_w, y0)
        self.cell(right_w, 5.6, pdf_text(right), align="R", new_x="LMARGIN", new_y="NEXT")

    def bullet(self, text: str) -> None:
        self.set_font("Times", "", 10.5)
        self.set_text_color(0, 0, 0)
        indent = self.l_margin + 5
        self.set_x(indent)
        self.multi_cell(self.w - indent - self.r_margin, 4.6, pdf_text(f"-  {text}"))

    def link_line(self, name_text: str, url: str) -> None:
        self.set_font("Times", "B", 11)
        self.set_text_color(*self.link_blue)
        self.write(5.6, pdf_text(name_text), link=url)
        self.set_text_color(0, 0, 0)
        self.ln(6)


def build_pdf() -> None:
    pdf = CVPdf()
    pdf.add_page()

    pdf.set_font("Times", "", 19)
    pdf.set_text_color(0, 0, 0)
    pdf.cell(0, 9, NAME, new_x="LMARGIN", new_y="NEXT", align="C")

    pdf.set_font("Times", "", 10.5)
    pdf.set_text_color(*pdf.link_blue)
    contact = f"{GITHUB}  |  {EMAIL}  |  {WEBSITE}"
    pdf.cell(0, 5.5, pdf_text(contact), new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.set_text_color(0, 0, 0)
    pdf.ln(1)

    # Experience
    pdf.section_heading("Experience")
    for exp in EXPERIENCE:
        pdf.exp_line(exp["title"], exp["org"], exp["dates"], org_url=exp.get("org_url"))
        for b in exp["bullets"]:
            pdf.bullet(b)
        pdf.ln(0.5)

    # Education
    pdf.section_heading("Education")
    for edu in EDUCATION:
        pdf.two_col_line(f"{edu['degree']} -- {edu['school']}", edu["dates"])
    pdf.ln(0.5)

    # Projects
    pdf.section_heading("Projects")
    for proj in PROJECTS:
        pdf.link_line(proj["name"], proj["url"])
        for b in proj["bullets"]:
            pdf.bullet(b)
        pdf.ln(0.5)

    # Open Source
    pdf.section_heading("Open Source")
    for proj in OPEN_SOURCE:
        pdf.link_line(proj["name"], proj["url"])
        for b in proj["bullets"]:
            pdf.bullet(b)
        pdf.ln(0.5)

    # Skills
    pdf.section_heading("Skills")
    for label, value in SKILLS:
        pdf.set_font("Times", "B", 10.5)
        pdf.write(5.6, pdf_text(f"{label}: "))
        pdf.set_font("Times", "", 10.5)
        pdf.write(5.6, pdf_text(value))
        pdf.ln(6.5)

    pdf.output(PDF_PATH)


def main() -> None:
    build_docx()
    build_pdf()
    print(f"Created: {DOCX_PATH}")
    print(f"Created: {PDF_PATH}")


if __name__ == "__main__":
    main()
